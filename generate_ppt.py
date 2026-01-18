#!/usr/bin/env python3
print("=== RUNNING LOCAL GENERATE_PPT.PY ===")
print("FILE:", __file__)
DEBUG = True

"""
PowerPoint Generator MVP

Generates professional PowerPoint presentations from JSON slide models.
"""

import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER


LAYOUT_INDICES = {
    "title": 0,
    "key_insights": 1,
    "chart": 2,
    "two_column": 3,
    "recommendation": 4,
    "appendix": 5,
}


def debug_print_placeholders(slide, slide_type: str):
    if not DEBUG:
        return
    print(f"\n[DEBUG] Placeholders for {slide_type}:")
    if not hasattr(slide, 'shapes') or not hasattr(slide.shapes, 'placeholders'):
        print("  No placeholders found")
        return
    for placeholder in slide.shapes.placeholders:
        try:
            idx = placeholder.placeholder_format.idx
            ph_type = placeholder.placeholder_format.type
            shape_type = placeholder.shape_type
            print(f"  idx={idx}, type={ph_type}, shape_type={shape_type}")
        except:
            pass


def get_title_placeholder(slide):
    for shape in slide.shapes:
        if not hasattr(shape, 'is_placeholder'):
            continue
        if shape.is_placeholder:
            try:
                if shape.placeholder_format.type == PP_PLACEHOLDER.TITLE:
                    return shape
            except:
                pass
    return None


def get_body_placeholder(slide):
    for shape in slide.shapes:
        if not hasattr(shape, 'is_placeholder'):
            continue
        if shape.is_placeholder:
            try:
                ph_type = shape.placeholder_format.type
                if ph_type in [PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT]:
                    return shape
            except:
                pass
    return None


def get_chart_placeholder(slide):
    for shape in slide.shapes:
        if not hasattr(shape, 'is_placeholder'):
            continue
        if shape.is_placeholder:
            try:
                if shape.placeholder_format.type == PP_PLACEHOLDER.CHART:
                    return shape
            except:
                pass
    for shape in slide.shapes:
        if not hasattr(shape, 'is_placeholder'):
            continue
        if shape.is_placeholder:
            try:
                if shape.placeholder_format.type == PP_PLACEHOLDER.OBJECT:
                    return shape
            except:
                pass
    return None


def list_all_placeholders(slide) -> str:
    if not hasattr(slide, 'shapes') or not hasattr(slide.shapes, 'placeholders'):
        return "No placeholders found"
    result = []
    for placeholder in slide.shapes.placeholders:
        try:
            idx = placeholder.placeholder_format.idx
            ph_type = placeholder.placeholder_format.type
            shape_type = placeholder.shape_type
            result.append(f"  idx={idx}, type={ph_type}, shape_type={shape_type}")
        except Exception as e:
            result.append(f"  Error: {e}")
    return "\n".join(result) if result else "No placeholders found"


def load_json_model(json_path: str) -> dict:
    with open(json_path, 'r', encoding='utf-8') as f:
        return json.load(f)


def create_title_slide(slide, slide_data: dict):
    debug_print_placeholders(slide, "title")
    title = slide_data.get("title", "")
    subtitle = slide_data.get("content", {}).get("subtitle", "")

    title_placeholder = get_title_placeholder(slide)
    if title_placeholder:
        title_placeholder.text = title
    else:
        raise ValueError(f"Title slide has no title placeholder.\n{list_all_placeholders(slide)}")

    subtitle_placeholder = get_body_placeholder(slide)
    if subtitle_placeholder:
        subtitle_placeholder.text = subtitle

    speaker_note = slide_data.get("speaker_note", "")
    if speaker_note:
        slide.notes_slide.notes_text_frame.text = speaker_note


def create_key_insights_slide(slide, slide_data: dict):
    debug_print_placeholders(slide, "key_insights")
    title = slide_data.get("title", "")
    bullets = slide_data.get("content", {}).get("bullets", [])

    title_placeholder = get_title_placeholder(slide)
    if title_placeholder:
        title_placeholder.text = title
    else:
        raise ValueError(f"Key insights slide has no title placeholder.\n{list_all_placeholders(slide)}")

    body_placeholder = get_body_placeholder(slide)
    if not body_placeholder:
        raise ValueError(f"Key insights slide has no body placeholder.\n{list_all_placeholders(slide)}")

    text_frame = body_placeholder.text_frame
    text_frame.clear()

    for idx, bullet in enumerate(bullets):
        if idx == 0:
            p = text_frame.paragraphs[0]
        else:
            p = text_frame.add_paragraph()
        p.text = bullet
        p.level = 0

    speaker_note = slide_data.get("speaker_note", "")
    if speaker_note:
        slide.notes_slide.notes_text_frame.text = speaker_note


def create_chart_slide(slide, slide_data: dict):
    """
    Create chart slide with storytelling enhancements.

    Enhancements:
    - Data labels (when requested)
    - Legend removal (when requested)
    - Chart enlargement (when requested)
    - Visual emphasis markers
    """
    debug_print_placeholders(slide, "chart")
    title = slide_data.get("title", "")
    content = slide_data.get("content", {})
    chart_type = content.get("chart_type", "column")
    data = content.get("data", {})

    title_placeholder = get_title_placeholder(slide)
    if title_placeholder:
        title_placeholder.text = title
    else:
        raise ValueError(f"Chart slide has no title placeholder.\n{list_all_placeholders(slide)}")

    categories = data.get("categories", [])
    series_list = data.get("series", [])

    if not categories:
        raise ValueError(f"Chart slide {slide_data.get('slide_number')}: missing categories")
    if not series_list:
        raise ValueError(f"Chart slide {slide_data.get('slide_number')}: missing series")

    expected_count = len(categories)
    for series in series_list:
        series_values = series.get("values", [])
        if len(series_values) != expected_count:
            raise ValueError(f"Chart slide {slide_data.get('slide_number')}: Series '{series.get('name')}' has {len(series_values)} values but should have {expected_count}")

    chart_data = ChartData()
    chart_data.categories = categories

    for series in series_list:
        series_name = series.get("name", "Series")
        series_values = series.get("values", [])
        chart_data.add_series(series_name, series_values)

    chart_type_map = {
        "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "line": XL_CHART_TYPE.LINE,
        "pie": XL_CHART_TYPE.PIE,
    }

    if chart_type not in chart_type_map:
        raise ValueError(f"Chart slide {slide_data.get('slide_number')}: Unsupported chart type '{chart_type}'")

    xl_chart_type = chart_type_map[chart_type]

    chart_placeholder = get_chart_placeholder(slide)
    if not chart_placeholder:
        raise ValueError(f"Chart slide has no chart placeholder.\n{list_all_placeholders(slide)}")

    # Check for visual emphasis (enlarge chart)
    visual_emphasis = content.get("visual_emphasis")
    if visual_emphasis == "large_chart":
        # Enlarge chart placeholder by 20%
        chart_placeholder.width = int(chart_placeholder.width * 1.2)
        chart_placeholder.height = int(chart_placeholder.height * 1.2)

    chart = None
    try:
        if DEBUG:
            print(f"[DEBUG] Attempting insert_chart on placeholder...")
        chart = chart_placeholder.insert_chart(xl_chart_type, chart_data)
        if DEBUG:
            print(f"[DEBUG] ✓ Chart inserted successfully")
    except AttributeError as e:
        if DEBUG:
            print(f"[DEBUG] insert_chart not supported, falling back to add_chart: {e}")
        try:
            x = chart_placeholder.left
            y = chart_placeholder.top
            cx = chart_placeholder.width
            cy = chart_placeholder.height
            chart = slide.shapes.add_chart(xl_chart_type, x, y, cx, cy, chart_data)
            if DEBUG:
                print(f"[DEBUG] ✓ Chart added via add_chart fallback")
        except Exception as fallback_error:
            raise ValueError(f"Failed to insert chart: insert_chart={e}, add_chart={fallback_error}")
    except Exception as e:
        raise ValueError(f"Failed to insert chart: {e}")

    if not chart:
        raise ValueError(f"Chart creation failed")

    # Apply chart storytelling enhancements
    try:
        chart_obj = chart.chart

        # Show/hide legend based on data
        show_legend = data.get("show_legend", True)
        if hasattr(chart_obj, 'has_legend'):
            chart_obj.has_legend = show_legend

        # Add data labels if requested
        show_data_labels = data.get("show_data_labels", False)
        if show_data_labels and hasattr(chart_obj, 'plots'):
            for plot in chart_obj.plots:
                plot.has_data_labels = True
                if hasattr(plot.data_labels, 'font'):
                    plot.data_labels.font.size = Pt(10)

        # Simplify gridlines (less clutter)
        if hasattr(chart_obj, 'value_axis'):
            value_axis = chart_obj.value_axis
            if hasattr(value_axis, 'has_major_gridlines'):
                value_axis.has_major_gridlines = True
            if hasattr(value_axis, 'has_minor_gridlines'):
                value_axis.has_minor_gridlines = False

    except Exception as e:
        # Chart enhancement failed, but chart exists - continue
        if DEBUG:
            print(f"[DEBUG] Chart enhancement warning: {e}")

    speaker_note = slide_data.get("speaker_note", "")
    if speaker_note:
        slide.notes_slide.notes_text_frame.text = speaker_note


def create_two_column_slide(slide, slide_data: dict):
    raise NotImplementedError("Two-column slide type not yet implemented")


def create_recommendation_slide(slide, slide_data: dict):
    """
    Recommendation slide layout:
    - Title
    - Recommendation (main text)
    - Rationale (bullets)
    - Next Steps (bullets)
    """
    debug_print_placeholders(slide, "recommendation")
    title = slide_data.get("title", "")
    content = slide_data.get("content", {})

    recommendation = content.get("recommendation", "")
    rationale = content.get("rationale", [])
    next_steps = content.get("next_steps", [])

    # Set title
    title_placeholder = get_title_placeholder(slide)
    if title_placeholder:
        title_placeholder.text = title
    else:
        raise ValueError(f"Recommendation slide has no title placeholder.\n{list_all_placeholders(slide)}")

    # Set body content
    body_placeholder = get_body_placeholder(slide)
    if not body_placeholder:
        raise ValueError(f"Recommendation slide has no body placeholder.\n{list_all_placeholders(slide)}")

    text_frame = body_placeholder.text_frame
    text_frame.clear()

    # Add recommendation as bold headline
    if recommendation:
        p = text_frame.paragraphs[0]
        p.text = recommendation
        p.level = 0
        # Make it bold for emphasis
        for run in p.runs:
            run.font.bold = True

    # Add rationale section
    if rationale:
        p = text_frame.add_paragraph()
        p.text = "Why:"
        p.level = 0
        for run in p.runs:
            run.font.bold = True

        for bullet in rationale:
            p = text_frame.add_paragraph()
            p.text = bullet
            p.level = 1

    # Add next steps section
    if next_steps:
        p = text_frame.add_paragraph()
        p.text = "Next Steps:"
        p.level = 0
        for run in p.runs:
            run.font.bold = True

        for step in next_steps:
            p = text_frame.add_paragraph()
            p.text = step
            p.level = 1

    # Add speaker note
    speaker_note = slide_data.get("speaker_note", "")
    if speaker_note:
        slide.notes_slide.notes_text_frame.text = speaker_note


def create_appendix_slide(slide, slide_data: dict):
    """
    Appendix slide layout:
    - Title
    - Items (bullet list)
    """
    debug_print_placeholders(slide, "appendix")
    title = slide_data.get("title", "")
    content = slide_data.get("content", {})
    items = content.get("items", [])

    # Set title
    title_placeholder = get_title_placeholder(slide)
    if title_placeholder:
        title_placeholder.text = title
    else:
        raise ValueError(f"Appendix slide has no title placeholder.\n{list_all_placeholders(slide)}")

    # Set body content
    body_placeholder = get_body_placeholder(slide)
    if not body_placeholder:
        raise ValueError(f"Appendix slide has no body placeholder.\n{list_all_placeholders(slide)}")

    text_frame = body_placeholder.text_frame
    text_frame.clear()

    # Add items as bullets
    for idx, item in enumerate(items):
        if idx == 0:
            p = text_frame.paragraphs[0]
        else:
            p = text_frame.add_paragraph()
        p.text = item
        p.level = 0

    # Add speaker note
    speaker_note = slide_data.get("speaker_note", "")
    if speaker_note:
        slide.notes_slide.notes_text_frame.text = speaker_note


SLIDE_HANDLERS = {
    "title": create_title_slide,
    "key_insights": create_key_insights_slide,
    "chart": create_chart_slide,
    "two_column": create_two_column_slide,
    "recommendation": create_recommendation_slide,
    "appendix": create_appendix_slide,
}


def validate_template(prs: Presentation):
    num_layouts = len(prs.slide_layouts)
    required_layouts = max(LAYOUT_INDICES.values()) + 1
    if num_layouts < required_layouts:
        raise ValueError(f"Template has only {num_layouts} layouts, need at least {required_layouts}")
    print(f"✓ Template validated: {num_layouts} layouts found")
    chart_layout_index = LAYOUT_INDICES.get("chart")
    if chart_layout_index is not None:
        chart_layout = prs.slide_layouts[chart_layout_index]
        print(f"  - Chart layout (index {chart_layout_index}): '{chart_layout.name}'")


def generate_presentation(template_path: str, json_path: str, output_path: str):
    prs = Presentation(template_path)
    validate_template(prs)

    data = load_json_model(json_path)
    presentation_data = data.get("presentation", {})
    slides_data = presentation_data.get("slides", [])

    for slide_data in slides_data:
        slide_type = slide_data.get("type")

        if slide_type not in LAYOUT_INDICES:
            print(f"✗ Unknown slide type '{slide_type}', skipping slide {slide_data.get('slide_number')}")
            continue

        if slide_type not in SLIDE_HANDLERS:
            print(f"✗ No handler for slide type '{slide_type}', skipping slide {slide_data.get('slide_number')}")
            continue

        layout_index = LAYOUT_INDICES[slide_type]
        slide_layout = prs.slide_layouts[layout_index]
        slide = prs.slides.add_slide(slide_layout)
        slide_index = len(prs.slides) - 1  # Track position for deletion if needed

        try:
            handler = SLIDE_HANDLERS[slide_type]
            handler(slide, slide_data)
            print(f"✓ Generated slide {slide_data.get('slide_number')}: {slide_type}")
        except NotImplementedError as e:
            # Remove the empty slide we just added
            rId = prs.slides._sldIdLst[slide_index].rId
            prs.part.drop_rel(rId)
            del prs.slides._sldIdLst[slide_index]
            print(f"✗ Removed slide {slide_data.get('slide_number')}: {e}")
        except Exception as e:
            # Remove the empty slide we just added
            try:
                rId = prs.slides._sldIdLst[slide_index].rId
                prs.part.drop_rel(rId)
                del prs.slides._sldIdLst[slide_index]
                print(f"✗ Removed slide {slide_data.get('slide_number')}: Rendering failed - {e}")
            except:
                print(f"✗ Error rendering slide {slide_data.get('slide_number')}: {e} (could not remove empty slide)")

    prs.save(output_path)
    print(f"\n✓ Presentation saved to: {output_path}")


if __name__ == "__main__":
    TEMPLATE_PATH = "template.pptx"
    # Prefer creative_output.json if it exists (from creative director)
    # Fall back to claude_output.json for direct rendering
    JSON_PATH = "creative_output.json" if Path("creative_output.json").exists() else "claude_output.json"
    OUTPUT_PATH = "output.pptx"

    if not Path(TEMPLATE_PATH).exists():
        print(f"Error: Template file not found: {TEMPLATE_PATH}")
        exit(1)

    if not Path(JSON_PATH).exists():
        print(f"Error: JSON file not found: {JSON_PATH}")
        exit(1)

    print("Starting PowerPoint generation...\n")
    generate_presentation(TEMPLATE_PATH, JSON_PATH, OUTPUT_PATH)
